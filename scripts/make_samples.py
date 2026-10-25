"""Generate the sample emails with PDF and Excel attachments in samples/inbox. All data is fictitious."""

import io
from email.message import EmailMessage
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook

INBOX = Path(__file__).resolve().parent.parent / "samples" / "inbox"


def make_pdf() -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Seetal Elektro AG - Bestellung SE-77120", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=10)
    for text in ("Datum: 21.10.2026", "Lieferadresse: Baustelle Hochdorf, Bahnhofstrasse 12, 6280 Hochdorf", ""):
        pdf.cell(0, 6, text, new_x="LMARGIN", new_y="NEXT")
    rows = [
        ("Pos", "Artikel", "Bezeichnung", "Menge", "Einheit", "Preis CHF"),
        ("1", "KAB-3x1.5", "Kabel 3x1.5", "2", "Rollen", "89.00"),
        ("2", "KAB-5X2.5-GR", "Kabel 5x2.5 grau", "150", "m", "3.40"),
        ("3", "ROHR-M20", "Rohr M20", "3", "Rolle", "45.00"),
        ("4", "WAGO 221-413", "Verbindungsklemme", "50", "Stk.", "0.55"),
    ]
    for row in rows:
        for width, value in zip((12, 35, 63, 20, 25, 25), row):
            pdf.cell(width, 7, value, border=1)
        pdf.ln()
    return bytes(pdf.output())


def make_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bestellung"
    ws.append(["Jura Haustechnik AG", None, "Bestellnr.", "JH-4501"])
    ws.append([])
    ws.append(["Artikel", "Bezeichnung", "Menge", "Einheit", "Preis CHF"])
    ws.append(["JH-KABEL-5", "Kabel 5x2.5", 2, "Rolle", 115.00])
    ws.append(["DOSE-AP-1", "Abzweigdose", 2.5, "Stk", 4.20])
    # No price on purpose: the extractor must return null, not invent one
    ws.append(["7610000000011", "Schraube M6x20", 400, "Stk", None])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def write_email(name: str, sender: str, subject: str, date: str, body: str,
                attachment: tuple[str, bytes, str]) -> None:
    message = EmailMessage()
    message["From"], message["To"] = sender, "bestellung@beispiel-technik.example"
    message["Subject"], message["Date"] = subject, date
    message.set_content(body)
    filename, data, subtype = attachment
    message.add_attachment(data, maintype="application", subtype=subtype, filename=filename)
    (INBOX / name).write_bytes(message.as_bytes())


if __name__ == "__main__":
    write_email("002-seetal.eml", "Beat Beispiel <einkauf@seetal-elektro.example>", "Bestellung SE-77120",
                "Wed, 21 Oct 2026 09:30:00 +0200", "Hallo\n\nAnbei unsere Bestellung.\n\nGruss\nBeat",
                ("SE-77120.pdf", make_pdf(), "pdf"))
    write_email("003-jura.eml", "Marco Beispiel <m.beispiel@jura-haustechnik.example>", "Bestellung JH-4501",
                "Wed, 21 Oct 2026 14:05:00 +0200", "Guten Tag\n\nBestellung im Anhang.\n\nMarco",
                ("Bestellung_JH-4501.xlsx", make_xlsx(),
                 "vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
