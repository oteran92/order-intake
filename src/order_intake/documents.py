"""Turn incoming files (email with attachments, PDF, Excel, CSV) into plain text for extraction."""

import io
from dataclasses import dataclass
from email import policy
from email.parser import BytesParser
from email.utils import parseaddr
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook


class UnreadableDocument(Exception):
    """The file has no text we can read, e.g. a scanned PDF (OCR is not supported yet)."""


@dataclass
class Document:
    source_file: str
    sender: str | None  # from the email header; None for files received without an email
    text: str


def read_document(path: Path) -> Document:
    data = path.read_bytes()
    if path.suffix.lower() == ".eml":
        return _read_email(path.name, data)
    return Document(path.name, None, _attachment_text(path.name, data))


def _read_email(name: str, data: bytes) -> Document:
    message = BytesParser(policy=policy.default).parsebytes(data)
    parts = [f"Subject: {message['subject'] or ''}", f"Date: {message['date'] or ''}"]
    body = message.get_body(preferencelist=("plain", "html"))
    if body is not None:
        parts.append(body.get_content())
    for attachment in message.iter_attachments():
        filename = attachment.get_filename() or "attachment"
        text = _attachment_text(filename, attachment.get_content())
        parts.append(f"--- Attachment: {filename} ---\n{text}")
    sender = parseaddr(message["from"] or "")[1] or None
    return Document(name, sender, "\n\n".join(parts))


def _attachment_text(filename: str, data: bytes | str) -> str:
    suffix = Path(filename).suffix.lower()
    if isinstance(data, str):
        return data
    if suffix == ".pdf":
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages).strip()
        if not text:
            raise UnreadableDocument(f"{filename} has no text layer")
        return text
    if suffix in (".xlsx", ".xlsm"):
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        # Tab-separated rows keep the table structure readable for the model
        return "\n".join(
            "\t".join("" if cell is None else str(cell) for cell in row)
            for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True)
        )
    if suffix in (".csv", ".txt"):
        return data.decode("utf-8", errors="replace")
    raise UnreadableDocument(f"{filename}: unsupported file type")
