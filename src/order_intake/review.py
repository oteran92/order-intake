"""Local review screen: the original document beside the extracted lines, with exceptions and approval."""

import io
import json
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from email import policy
from email.parser import BytesParser
from pathlib import Path
from urllib.parse import urlencode

from fastapi import FastAPI, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook
from pydantic import BaseModel

from .catalog import Catalog
from .export import export_document, to_csv, to_json
from .models import Correction, ExtractedOrder, Severity, ValidationResult
from .validate import DuplicateIndex, validate_order

TEMPLATES = Jinja2Templates(directory=Path(__file__).parent / "templates")


@dataclass
class SourcePart:
    """One readable piece of the original document: email text, a PDF or a spreadsheet."""

    title: str
    kind: str  # "text", "pdf" or "table"
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)
    data: bytes = b""


@dataclass
class ReviewItem:
    stem: str
    order: ExtractedOrder
    result: ValidationResult
    source: Path

    def has(self, severity: Severity) -> bool:
        return any(issue.severity is severity for issue in self.result.issues)


def source_parts(path: Path) -> list[SourcePart]:
    data = path.read_bytes()
    if path.suffix.lower() != ".eml":
        return [_part(path.name, data)]
    message = BytesParser(policy=policy.default).parsebytes(data)
    header = "\n".join(f"{name}: {message[name]}" for name in ("From", "To", "Date", "Subject") if message[name])
    body = message.get_body(preferencelist=("plain",))
    parts = [SourcePart("Email", "text", text=f"{header}\n\n{body.get_content() if body else ''}")]
    parts += [_part(a.get_filename() or "attachment", a.get_content()) for a in message.iter_attachments()]
    return parts


def _part(filename: str, data: bytes | str) -> SourcePart:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return SourcePart(filename, "pdf", data=data)
    if suffix in (".xlsx", ".xlsm"):
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = [["" if cell is None else str(cell) for cell in row]
                for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True)]
        return SourcePart(filename, "table", rows=rows)
    text = data if isinstance(data, str) else data.decode("utf-8", errors="replace")
    return SourcePart(filename, "text", text=text)


def approval_record(item: ReviewItem) -> dict:
    """The approved order handed to the ERP export: base-unit quantities plus the line as ordered."""
    header = item.order.model_dump(mode="json", exclude={"lines"})
    lines = [
        {"line_no": validated.line_no, "item_no": validated.item_no, "quantity": validated.quantity,
         "unit": validated.unit, "order_quantity": raw.quantity, "order_unit": raw.unit,
         "unit_price": raw.unit_price}
        for raw, validated in zip(item.order.lines, item.result.lines)
    ]
    confirmed = sorted({i.code for i in item.result.issues if i.severity is Severity.WARNING})
    return header | {"customer_id": item.result.customer_id, "lines": lines, "confirmed_warnings": confirmed,
                     "approved_at": datetime.now(timezone.utc).isoformat(timespec="seconds")}


def rejection_record(item: ReviewItem, reason: str | None) -> dict:
    """The order as it stood when rejected, with its corrections and open issues, so the decision can be traced."""
    return item.order.model_dump(mode="json") | {
        "customer_id": item.result.customer_id,
        "validated_lines": [line.model_dump(mode="json") for line in item.result.lines],
        "open_issues": [issue.model_dump(mode="json") for issue in item.result.issues],
        "reason": reason,
        "rejected_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
    }


class InvalidInput(ValueError):
    pass


def parse_number(value: str, name: str, required: bool = False) -> float | None:
    """Accept Swiss and German notation as typed by a reviewer: 1'000, 2,5."""
    cleaned = value.strip().replace("'", "").replace(",", ".")
    if not cleaned:
        if required:
            raise InvalidInput(f"{name} is required")
        return None
    try:
        return float(cleaned)
    except ValueError:
        raise InvalidInput(f"{name}: '{value}' is not a number") from None


def parse_date(value: str, name: str) -> date | None:
    if not value.strip():
        return None
    try:
        return date.fromisoformat(value.strip())
    except ValueError:
        raise InvalidInput(f"{name}: '{value}' is not a date (YYYY-MM-DD)") from None


def apply_changes(order: ExtractedOrder, target: BaseModel, values: dict, line_no: int | None) -> bool:
    """Set changed fields on the order or one of its lines and log each change. Returns True if anything changed."""
    now = datetime.now(timezone.utc)
    changed = False
    for name, value in values.items():
        old = getattr(target, name)
        if old != value:
            order.corrections.append(Correction(line_no=line_no, field=name, at=now,
                                                old=None if old is None else str(old),
                                                new=None if value is None else str(value)))
            setattr(target, name, value)
            changed = True
    return changed


def load_items(orders_dir: Path, inbox: Path, catalog_dir: Path) -> dict[str, ReviewItem]:
    catalog = Catalog.from_dir(catalog_dir)
    duplicates = DuplicateIndex()
    items = {}
    # Sorted by sequence number so duplicate detection follows arrival order, as in the CLI
    for path in sorted(orders_dir.glob("*.json")):
        order = ExtractedOrder.model_validate_json(path.read_text(encoding="utf-8"))
        result = validate_order(order, catalog, duplicates)
        items[path.stem] = ReviewItem(path.stem, order, result, inbox / order.source_file)
    return items


def create_app(orders_dir: Path, inbox: Path, catalog_dir: Path, approved_dir: Path,
               rejected_dir: Path | None = None) -> FastAPI:
    app = FastAPI(title="order-intake review")
    items = load_items(orders_dir, inbox, catalog_dir)
    rejected_dir = rejected_dir or approved_dir.with_name("rejected")
    decision_dirs = {"approved": approved_dir, "rejected": rejected_dir}
    for folder in decision_dirs.values():
        folder.mkdir(parents=True, exist_ok=True)

    def get_item(stem: str) -> ReviewItem:
        if stem not in items:
            raise HTTPException(404, "Unknown order")
        return items[stem]

    def decision(stem: str) -> str | None:
        """'approved', 'rejected' or None while the order is open. Both decisions are final."""
        return next((name for name, folder in decision_dirs.items() if (folder / f"{stem}.json").exists()), None)

    def rejection_reason(stem: str) -> str | None:
        path = rejected_dir / f"{stem}.json"
        return json.loads(path.read_text(encoding="utf-8")).get("reason") if path.exists() else None

    def save_and_revalidate(item: ReviewItem) -> None:
        (orders_dir / f"{item.stem}.json").write_text(item.order.model_dump_json(indent=2), encoding="utf-8")
        # Revalidate everything: a changed PO number or line can create or clear a duplicate elsewhere
        items.clear()
        items.update(load_items(orders_dir, inbox, catalog_dir))

    def edit(stem: str, change) -> RedirectResponse:
        """Shared flow for header and line edits: refuse decided orders, apply, save, show the result."""
        item = get_item(stem)
        if state := decision(stem):
            raise HTTPException(409, f"{state.capitalize()} orders cannot be edited")
        try:
            changed = change(item)
        except InvalidInput as error:
            return RedirectResponse(f"/orders/{stem}?{urlencode({'error': str(error)})}", status_code=303)
        if changed:
            save_and_revalidate(item)
        return RedirectResponse(f"/orders/{stem}?notice={'saved' if changed else 'unchanged'}", status_code=303)

    @app.post("/orders/{stem}/header")
    def edit_header(stem: str, po_number: str = Form(""), order_date: str = Form(""),
                    delivery_date: str = Form(""), ship_to: str = Form("")):
        def change(item: ReviewItem) -> bool:
            values = {"po_number": po_number.strip() or None, "order_date": parse_date(order_date, "Order date"),
                      "delivery_date": parse_date(delivery_date, "Delivery date"), "ship_to": ship_to.strip() or None}
            return apply_changes(item.order, item.order, values, None)
        return edit(stem, change)

    @app.post("/orders/{stem}/lines/{line_no}")
    def edit_line(stem: str, line_no: int, customer_ref: str = Form(""), description: str = Form(""),
                  quantity: str = Form(""), unit: str = Form(""), unit_price: str = Form("")):
        def change(item: ReviewItem) -> bool:
            line = next((l for l in item.order.lines if l.line_no == line_no), None)
            if line is None:
                raise HTTPException(404, "Unknown line")
            values = {"customer_ref": customer_ref.strip(), "description": description.strip() or None,
                      "quantity": parse_number(quantity, "Quantity", required=True), "unit": unit.strip() or None,
                      "unit_price": parse_number(unit_price, "Price")}
            return apply_changes(item.order, line, values, line_no)
        return edit(stem, change)

    @app.get("/", response_class=HTMLResponse)
    def index(request: Request):
        rows = [(item, decision(item.stem)) for item in items.values()]
        return TEMPLATES.TemplateResponse(request, "index.html", {"rows": rows})

    @app.get("/orders/{stem}", response_class=HTMLResponse)
    def order_page(request: Request, stem: str, notice: str | None = None, error: str | None = None):
        item = get_item(stem)
        by_line: dict[int, list] = {}
        for issue in item.result.issues:
            by_line.setdefault(issue.line_no, []).append(issue)
        parts = source_parts(item.source) if item.source.exists() else []
        context = {
            "item": item, "parts": parts, "lines": list(zip(item.order.lines, item.result.lines)),
            "header_issues": by_line.get(None, []), "line_issues": by_line,
            "decision": decision(stem), "rejection_reason": rejection_reason(stem),
            "has_errors": item.has(Severity.ERROR), "has_warnings": item.has(Severity.WARNING), "notice": notice, "error": error,
        }
        return TEMPLATES.TemplateResponse(request, "order.html", context)

    @app.get("/orders/{stem}/parts/{index}")
    def part_file(stem: str, index: int):
        item = get_item(stem)
        parts = source_parts(item.source)
        if not 0 <= index < len(parts) or parts[index].kind != "pdf":
            raise HTTPException(404, "No PDF at this position")
        return Response(parts[index].data, media_type="application/pdf",
                        headers={"Content-Disposition": f'inline; filename="{parts[index].title}"'})

    @app.post("/orders/{stem}/approve")
    def approve(stem: str, confirm_warnings: str | None = Form(None)):
        item = get_item(stem)
        if decision(stem) == "rejected":
            raise HTTPException(409, "Rejected orders cannot be approved")
        if item.has(Severity.ERROR):
            raise HTTPException(409, "Errors must be resolved before approval")
        if item.has(Severity.WARNING) and not confirm_warnings:
            raise HTTPException(400, "Confirm the warnings before approval")
        target = approved_dir / f"{stem}.json"
        if not target.exists():
            target.write_text(json.dumps(approval_record(item), indent=2, ensure_ascii=False), encoding="utf-8")
        return RedirectResponse(f"/orders/{stem}?notice=approved", status_code=303)

    @app.get("/orders/{stem}/export.{fmt}")
    def download_export(stem: str, fmt: str):
        get_item(stem)
        path = approved_dir / f"{stem}.json"
        if fmt not in ("json", "csv") or not path.exists():
            raise HTTPException(404, "Only approved orders can be exported")
        document = export_document(json.loads(path.read_text(encoding="utf-8")), stem)
        body, media = (to_json(document), "application/json") if fmt == "json" else (to_csv(document), "text/csv")
        return Response(body, media_type=f"{media}; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="{stem}.{fmt}"'})

    @app.post("/orders/{stem}/reject")
    def reject(stem: str, reason: str = Form("")):
        item = get_item(stem)
        if decision(stem) == "approved":
            raise HTTPException(409, "Approved orders cannot be rejected")
        # Errors do not block rejection: an unusable order is the main reason to reject
        target = rejected_dir / f"{stem}.json"
        if not target.exists():
            record = rejection_record(item, reason.strip() or None)
            target.write_text(json.dumps(record, indent=2, ensure_ascii=False), encoding="utf-8")
        return RedirectResponse(f"/orders/{stem}?notice=rejected", status_code=303)

    return app
